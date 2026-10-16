# 仮の値だった全国・都道府県・市区町村の数字を、公表資料から取り出して data-src/stats/extra.json に保存する。
# 使い方：python3 -I tools/extra.py
#   資料は raw/extra に保存し、2回目からは保存したものを使う。e-Stat の表は環境変数 ICHIOKU_ESTAT_APPID のアプリIDで取る。
#   取り出した数字は、資料にある合計や別の表の値と照らし合わせ、合わなければ止まる。
import json, os, re, sys, unicodedata, urllib.request, urllib.parse, hashlib

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, 'raw', 'extra')
OUT = os.path.join(ROOT, 'data-src', 'stats', 'extra.json')
STATS = os.path.join(ROOT, 'data-src', 'stats', 'stats.json')
os.makedirs(CACHE, exist_ok=True)

PREFS = ['北海道','青森県','岩手県','宮城県','秋田県','山形県','福島県','茨城県','栃木県','群馬県','埼玉県','千葉県','東京都','神奈川県','新潟県','富山県','石川県','福井県','山梨県','長野県','岐阜県','静岡県','愛知県','三重県','滋賀県','京都府','大阪府','兵庫県','奈良県','和歌山県','鳥取県','島根県','岡山県','広島県','山口県','徳島県','香川県','愛媛県','高知県','福岡県','佐賀県','長崎県','熊本県','大分県','宮崎県','鹿児島県','沖縄県']
PCODE = {p: '%02d' % (i + 1) for i, p in enumerate(PREFS)}
# 部首の形の文字（⻘・⻑ など）は NFKC で直らないため、ふつうの漢字に置き換える
RAD = str.maketrans({'⻘': '青', '⻑': '長', '⻄': '西', '⻝': '食', '⻤': '鬼', '⻩': '黄', '⻆': '角', '⻌': '辶'})
N = lambda s: unicodedata.normalize('NFKC', s or '').translate(RAD)
problems = []


def fetch(url, name):
    p = os.path.join(CACHE, name)
    if not os.path.exists(p):
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=120) as r, open(p, 'wb') as f: f.write(r.read())
    with open(p, 'rb') as f:
        head = f.read(8)
    if name.endswith('.pdf') and not head.startswith(b'%PDF'): sys.exit(name + '：PDF ではないものが届いた（' + url + '）')
    return p


def estat(sid, params):
    key = hashlib.md5((sid + params).encode()).hexdigest()[:10]
    p = os.path.join(CACHE, 'estat-' + sid + '-' + key + '.json')
    if not os.path.exists(p):
        app = os.environ.get('ICHIOKU_ESTAT_APPID')
        if not app: sys.exit('環境変数 ICHIOKU_ESTAT_APPID がない')
        url = 'https://api.e-stat.go.jp/rest/3.0/app/json/getStatsData?appId=' + app + '&lang=J&metaGetFlg=Y&limit=100000&statsDataId=' + sid + '&' + params
        with urllib.request.urlopen(url, timeout=120) as r: j = json.loads(r.read())
        d = j['GET_STATS_DATA'].get('STATISTICAL_DATA')
        if not d or 'DATA_INF' not in d: sys.exit('e-Stat の表 ' + sid + ' が取れない：' + json.dumps(j['GET_STATS_DATA']['RESULT'], ensure_ascii=False))
        if d['RESULT_INF'].get('NEXT_KEY'): sys.exit('e-Stat の表 ' + sid + ' が10万件を超えた')
        with open(p, 'w') as f: json.dump(d, f, ensure_ascii=False)
    with open(p) as f: d = json.load(f)
    vals = d['DATA_INF']['VALUE']
    return vals if isinstance(vals, list) else [vals]


def pdf_pages(path):
    """PDF のページごとの文字。取り出しに時間がかかるため、2回目からは保存した文字を使う"""
    tp = path + '.pages.json'
    if os.path.exists(tp) and os.path.getmtime(tp) >= os.path.getmtime(path):
        with open(tp) as f: return [N(x) for x in json.load(f)]
    import pdfplumber
    with pdfplumber.open(path) as p: raw = [pg.extract_text() or '' for pg in p.pages]
    with open(tp, 'w') as f: json.dump(raw, f, ensure_ascii=False)
    return [N(x) for x in raw]


def need(cond, msg):
    if not cond: sys.exit('照らし合わせが合わない：' + msg)


out = {'sources': {}, 'nat': {}, 'prefs': {}, 'munis': {}, 'goods': {}}
fv = lambda v: float(v['$']) if re.fullmatch(r'-?[\d.]+', v.get('$', '')) else None

# ---------- 共働きの割合（労働力調査特別調査・詳細集計。労働政策研究・研修機構のまとめ表） ----------
import openpyxl
wb = openpyxl.load_workbook(fetch('https://www.jil.go.jp/kokunai/statistics/timeseries/xls/g0212.xlsx', 'dual-jil.xlsx'), data_only=True)
dual = {}
for r in wb.worksheets[0].iter_rows(values_only=True):
    y = r[1] if isinstance(r[1], int) else (int(r[1][:4]) if isinstance(r[1], str) and re.match(r'\d{4}年', r[1]) else None)
    if y and isinstance(r[2], (int, float)) and isinstance(r[3], (int, float)): dual[y] = round(r[3] / (r[2] + r[3]), 3)
need(len(dual) >= 45 and dual[1980] == round(614 / (1114 + 614), 3), '共働きの表（1980年は専業主婦1114万・共働き614万世帯）')
out['nat']['dualIncome'] = dual
out['sources']['dualIncome'] = '労働力調査特別調査（1980〜2001年）・労働力調査 詳細集計（2002年〜）。共働き世帯 ÷（共働き世帯＋専業主婦世帯）。1979年以前は1980年の値を使う'

# ---------- 50歳までに結婚する割合（国勢調査2020年の実績と、日本の世帯数の将来推計2024年の配偶関係別人口） ----------
wb = openpyxl.load_workbook(fetch('https://www.ipss.go.jp/pp-ajsetai/j/HPRJ2024/hprj2024_gaiyo_siryou1.xlsx', 'ipss-hh2024-siryou1.xlsx'), data_only=True)
em = {'m': {}, 'f': {}}
for ws in wb.worksheets:
    m = re.search(r'(\d{4})', ws.title)
    if not m: continue
    y, sex, got = int(m.group(1)), None, {}
    for r in ws.iter_rows(values_only=True):
        if r[0] in ('男', '女'): sex = 'm' if r[0] == '男' else 'f'
        if sex and isinstance(r[0], str) and r[0].replace(' ', '') in ('45～49歳', '50～54歳'): got[(sex, r[0].replace(' ', ''))] = r[6]
    for s in ('m', 'f'):
        u = (got[(s, '45～49歳')] + got[(s, '50～54歳')]) / 2 / 100
        em[s][y - 50] = round(1 - u, 4)
# 2020年は人口統計資料集の50歳時の未婚割合（男28.25%・女17.81%）と合うはず
need(abs(em['m'][1970] - (1 - 0.2825)) < 0.001 and abs(em['f'][1970] - (1 - 0.1781)) < 0.001, '50歳時の未婚割合（2020年）')
need(len(em['m']) == 7, '将来推計の年（2020〜2050年の7つ）')
out['nat']['everMarried'] = em
out['sources']['everMarried'] = '国勢調査（2020年）と日本の世帯数の将来推計（全国推計・2024年）の配偶関係別人口。その年に45〜49歳と50〜54歳の人のうち一度も結婚していない割合の平均を、50年前に生まれた人の値とした'

# ---------- 出産で仕事をやめる割合（第16回出生動向基本調査 図表9-1-3） ----------
t = '\n'.join(pdf_pages(fetch('https://www.ipss.go.jp/ps-doukou/j/doukou16/JNFS16_Report04.pdf', 'ipss-jnfs16-part2.pdf')))
i = t.index('図表9-1-3'); blk = t[i:t.index('第2子', i)]
leave = {}
for m in re.finditer(r'(\d{4})[～~](\d{2})年 ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)', blk):
    a, b, c = float(m.group(3)), float(m.group(4)), float(m.group(5))
    need(abs(sum(float(m.group(k)) for k in range(3, 8)) - 100) < 0.3, '第1子の就業変化の合計が100%（' + m.group(1) + '年）')
    leave[int(m.group(1)) + 2] = round(c / (a + b + c), 3)
need(len(leave) == 7 and abs((42.6 + 11.2) - 53.8) < 0.01, '第1子の出生年の7区分')
out['nat']['leaveAtBirth'] = leave
out['sources']['leaveAtBirth'] = '第16回出生動向基本調査（2021年）図表9-1-3。妊娠がわかったときに働いていた妻のうち、第1子が1歳のときに仕事をやめていた割合（出産退職 ÷（就業継続＋出産退職））。第1子の生まれ年5年ごとの真ん中の年に置いた'

# ---------- 2027年以降の平均寿命（日本の将来推計人口・令和5年推計 仮定値表6 死亡中位） ----------
pages = pdf_pages(fetch('https://www.ipss.go.jp/pp-zenkoku/j/zenkoku2023/pp2023_Report3.pdf', 'ipss-pp2023-report3.pdf'))
pg = next(p for p in pages if p.startswith('表6 男女別平均寿命'))
e0p = {'m': {}, 'f': {}}
for m in re.finditer(r'\((\d{4})\) ([\d.]+) ([\d.]+) ([\d.]+)', pg):
    y = int(m.group(1)); e0p['m'][y] = float(m.group(2)); e0p['f'][y] = float(m.group(3))
    need(abs(float(m.group(3)) - float(m.group(2)) - float(m.group(4))) < 0.015, '平均寿命の男女差（' + str(y) + '年）')
need(len(e0p['m']) == 51 and e0p['m'][2070] == 85.89 and e0p['f'][2070] == 91.94, '平均寿命の推計（2020〜2070年）')
out['nat']['e0proj'] = e0p
out['sources']['e0proj'] = '日本の将来推計人口（令和5年推計）仮定値表 表6（死亡中位）。2071年以降は2070年の値を使う'

# ---------- 子どもの時期の引っ越し（国勢調査：5年前に住んでいた場所 × 10〜14歳） ----------
def mv(sid, params, pick):
    vals = estat(sid, params)
    return pick(vals)
def sumv(vals, **f):
    s = 0; n = 0
    for v in vals:
        if all(v.get('@' + k) in (w if isinstance(w, (list, tuple)) else [w]) for k, w in f.items()):
            x = fv(v)
            if x is not None: s += x; n += 1
    return s if n else None
move = {}
v = estat('0000031780', 'cdCat02=000&cdCat03=202&cdArea=00000')
p5 = (sumv(v, cat01='07') + sumv(v, cat01='08')) / (sumv(v, cat01='02') + sumv(v, cat01='03') + sumv(v, cat01='14')); move[1990] = p5
v = estat('0000033389', 'cdCat01=000&cdCat03=202&cdArea=00000')
p5 = (sumv(v, cat02='004') + sumv(v, cat02='005')) / (sumv(v, cat02='001') + sumv(v, cat02='002') + sumv(v, cat02='007')); move[2000] = p5
v = estat('0003046138', 'cdCat03=000&cdArea=00000&cdCat02=011,012,013,014,015')
p5 = (sumv(v, cat01='00505') + sumv(v, cat01='00506')) / (sumv(v, cat01='00501') + sumv(v, cat01='00502') + sumv(v, cat01='00513')); move[2010] = p5
v = estat('0003447398', 'cdCat01=0&cdCat02=03&cdArea=00000')
p5 = (sumv(v, cat03='002112') + sumv(v, cat03='00212')) / (sumv(v, cat03='001') + sumv(v, cat03='0021') + sumv(v, cat03='0022')); move[2020] = p5
for y, p in move.items(): need(0.02 < p < 0.2, '10〜14歳の県内の引っ越しの割合（' + str(y) + '年 ' + str(round(p, 4)) + '）')
out['nat']['moveChild'] = {y - 3: round(1 - (1 - p) ** 0.2, 4) for y, p in move.items()}
out['nat']['moveChild5'] = {y: round(p, 4) for y, p in move.items()}
out['sources']['moveChild'] = '国勢調査（1990・2000・2010・2020年）の「5年前に住んでいた場所」。10〜14歳のうち、5年前は同じ県の別の市区町村（同じ市の別の区を含む）に住んでいた子の割合を、1年あたりに直した値。調査の3年前の年に置いた'

# ---------- 県外に出た人が地元へ戻る割合（第9回人口移動調査 図V-1） ----------
pages = pdf_pages(fetch('https://www.ipss.go.jp/ps-idou/j/migration/m09/ido9_hokoku.pdf', 'ipss-mig9-report.pdf'))
pg = next(p for p in pages if '図 V-1 出生都道府県別、県外移動歴' in p or '図V-1 出生都道府県別、県外移動歴' in p)
toks = pg.split()
ret = {}
names = set(PREFS) | {'全国'}
for i, w in enumerate(toks):
    if w not in names: continue
    nums = []
    for x in toks[i + 1:]:
        if re.fullmatch(r'\d+\.\d', x): nums.append(float(x))
        else: break
    if len(nums) == 4:
        need(abs(sum(nums) - 100) < 0.3, '県外移動歴の合計が100%（' + w + '）')
        need(w not in ret, '図V-1 に同じ県が2回（' + w + '）')
        ret[w] = round(nums[1] / (nums[1] + nums[2]), 3)
need(len(ret) == 48, '図V-1 の47都道府県と全国（' + str(len(ret)) + '）')
need(ret['北海道'] == round(17.5 / (17.5 + 19.8), 3), '北海道の値')
out['nat']['retNat'] = ret.pop('全国')
for p, r in ret.items(): out['prefs'].setdefault(PCODE[p], {})['ret'] = {2023: r}
out['sources']['ret'] = '第9回人口移動調査（2023年）図V-1。生まれた都道府県から県外へ出たことのある人のうち、生まれた都道府県に戻って住んでいる人（Uターン）の割合'

# ---------- 市区町村の一覧（照らし合わせ用） ----------
with open(STATS) as f: st = json.load(f)
munis = {}
for p in st['prefs']:
    for m in p['munis']: munis[m['code']] = (p['code'], m['name'])
# 資料によって「ケ・ヶ」「鰺・鯵」の書き方がちがうため、そろえてから比べる
K = lambda n: n.replace('ヶ', 'ケ').replace('鯵', '鰺').replace('檮', '梼')
byName = {}
for c, (pc, n) in munis.items(): byName.setdefault((pc, K(n)), []).append(c)
def cityCodes(pc, n):
    """名前の市区町村の番号。政令指定都市の名前なら、その区すべて"""
    n = K(n)
    if (pc, n) in byName: return byName[(pc, n)]
    return [c for c, (q, m) in munis.items() if q == pc and re.fullmatch(re.escape(n) + r'.+区', K(m))]

# ---------- 雪の印（豪雪地帯の指定市町村。国土交通省） ----------
pages = pdf_pages(fetch('https://www.mlit.go.jp/common/001475871.pdf', 'mlit-gousetsu.pdf'))
want = {}
for m in re.finditer(r'(' + '|'.join(PREFS) + r')\s*※?\s*(\d+)\s+(\d+)\s+\d+\s+\d+\s+\d+', pages[0]): want[m.group(1)] = int(m.group(3))
need(sum(want.values()) == 532 and len(want) == 24, '豪雪地帯の県別の市町村数（24道府県・計532）')
body = '\n'.join(pages[1:])
body = re.sub(r'\([^)]*\)', ' ', body)
snow, cur, cnt = set(), None, {}
for line in body.split('\n'):
    m = re.match(r'○(' + '|'.join(PREFS) + ')', line)
    if m: cur = m.group(1); cnt[cur] = 0; continue
    if not cur or re.match(r'(豪雪地帯|うち特別)', line.strip()): continue
    for w in re.sub(r'[●▲※]', ' ', line).split():
        if w.endswith('郡') or w.startswith('旧') or not re.search(r'[市町村]$', w) or w in ('市', '町', '村', '郡名'): continue
        cs = cityCodes(PCODE[cur], w)
        if not cs: problems.append('豪雪地帯の「' + cur + w + '」に合う市区町村がない'); continue
        cnt[cur] += 1; snow.update(cs)
bad = {p: (cnt.get(p), n) for p, n in want.items() if cnt.get(p) != n}
need(not bad, '豪雪地帯の県ごとの数（取り出した数, 資料の数）' + json.dumps(bad, ensure_ascii=False))

# ---------- 海沿いの印（2023年漁業センサス 海面漁業調査：海で漁をする経営体のある市区町村） ----------
p = os.path.join(CACHE, 'estat-meta-0002144718.json')
if not os.path.exists(p):
    url = 'https://api.e-stat.go.jp/rest/3.0/app/json/getMetaInfo?appId=' + os.environ['ICHIOKU_ESTAT_APPID'] + '&lang=J&statsDataId=0002144718'
    with urllib.request.urlopen(url, timeout=120) as r: j = json.loads(r.read())
    with open(p, 'w') as f: json.dump(j, f, ensure_ascii=False)
with open(p) as f: j = json.load(f)
cls = next(o for o in j['GET_META_INFO']['METADATA_INF']['CLASS_INF']['CLASS_OBJ'] if o['@id'] == 'cat01')['CLASS']
coast, wardCity = set(), set()
for c in cls:
    parts = c['@name'].split('_')
    if parts[0] not in PCODE: continue
    pc = PCODE[parts[0]]
    for w in parts[1:]:
        if (pc, K(w)) in byName: coast.update(byName[(pc, K(w))]); break
        if cityCodes(pc, w): wardCity.add(parts[0] + w); break
need(len(coast) > 400, '海沿いの市区町村の数（' + str(len(coast)) + '）')
for c in list(snow) + list(coast): need(c in munis, '印の番号 ' + c)
for c in munis:
    f = {}
    if c in snow: f['snow'] = True
    if c in coast: f['coast'] = True
    if f: out['munis'][c] = f
out['sources']['snow'] = '国土交通省「豪雪地帯及び特別豪雪地帯の指定地域」。一部の地域だけが指定の市町村も含む。政令指定都市はすべての区に付けた'
out['sources']['coast'] = '2023年漁業センサス 海面漁業調査（市区町村別の漁業経営体数）に出てくる市区町村。政令指定都市は区ごとの数字がないため印を付けていない（' + '、'.join(sorted(wardCity)) + '）'

# ---------- 1995年より前の市区町村の生まれる割合の手がかり（社会・人口統計体系：15歳未満人口。今の市区町村の区切りに組み直し済み） ----------
u14 = {}
for v in estat('0000020201', 'cdCat01=A1301'):
    x = fv(v)
    if x is None: continue
    u14.setdefault(v['@area'], {})[int(v['@time'][:4])] = x
miss = [c for c in munis if c not in u14 or 1980 not in u14[c]]
out['u14'] = {c: u14[c] for c in munis if c in u14}
out['u14Missing'] = miss
out['sources']['u14'] = '社会・人口統計体系 市区町村データ A1301（15歳未満人口。国勢調査、今の市区町村の区切りに組み直したもの）'

# ---------- 普及率の古い年（通信利用動向調査の報告書の本文と e-Stat の表） ----------
def sent(y, pat):
    t = re.sub(r'\s+', '', ''.join(pdf_pages(fetch('https://www.soumu.go.jp/johotsusintokei/statistics/pdf/HR%d00_001.pdf' % y, 'soumu-comm-%d.pdf' % y))))
    m = re.search(pat, t)
    need(m, '通信利用動向調査 %d年の報告書に「%s」' % (y, pat))
    return float(m.group(1))
mob = {1996: sent(1996, r'保有率は、携帯電話([\d.]+)%'), 1997: sent(1997, r'保有率は、携帯電話([\d.]+)%'), 1998: sent(1998, r'携帯電話はほぼ5世帯に3世帯の割合\(([\d.]+)%\)'),
       1999: sent(1999, r'携帯電話はほぼ3世帯に2世帯の割合\(([\d.]+)%\)'), 2000: sent(2000, r'携帯電話を保有する世帯の割合\(保有率\)はすでに7割\(([\d.]+)%\)'),
       2001: sent(2001, r'携帯電話を保有する世帯の割合\(保有率\)は([\d.]+)%')}
t2000 = re.sub(r'\s+', '', ''.join(pdf_pages(os.path.join(CACHE, 'soumu-comm-2000.pdf'))))
net = {1996: sent(1996, r'自宅でインターネットを利用している世帯\(利用率\)は([\d.]+)%'), 1997: sent(1997, r'インターネットの利用率は([\d.]+)%'),
       2000: sent(2000, r'インターネットを利用している世帯は、全世帯の3割強\(([\d.]+)%\)'), 2001: sent(2001, r'インターネットを利用している世帯は、全世帯の約6割\(([\d.]+)%\)'),
       2002: sent(2002, r'インターネットを利用している世帯は、全世帯の約8割\(([\d.]+)%\)'), 2003: sent(2003, r'インターネットを利用している世帯は9割弱\(([\d.]+)%\)'),
       2004: sent(2004, r'インターネットを利用している世帯は昨年より[\d.]+ポイント減少して([\d.]+)%'), 2005: sent(2005, r'インターネットを利用している世帯は昨年より[\d.]+ポイント増加して([\d.]+)%')}
# 1998・1999年は2000年の報告書の図表4-1（利用している・…の順の帯グラフの数字）から
blk = t2000[t2000.index('平成9年6.4'):]
net[1998] = float(re.search(r'平成10年(\d+\.\d)', blk).group(1))
net[1999] = float(re.search(r'平成11年(\d+\.\d)', blk).group(1))
need(net[1998] == 11.0 and net[1999] == 19.1, '1998・1999年のインターネットの利用率')
for v in estat('0003161383', 'cdTab=00100&cdCat04=100&cdCat01=110'): net[int(v['@time'][:4])] = fv(v)
smart = {int(v['@time'][:4]): fv(v) for v in estat('0003161336', 'cdTab=00100&cdCat04=100&cdCat01=260')}
need(smart.get(2010) == 9.7 and net.get(2006) == 89.3 and net[2002] == 81.4, '通信利用動向調査の e-Stat の表')
def ramp(y0, y1, v1):
    return {y: round(v1 * (y - y0) / (y1 - y0), 1) for y in range(y0, y1)}
out['goods']['mobile'] = {'real': mob, 'est': ramp(1987, 1996, mob[1996]), 'src': 'comm', 'note': '1996〜2001年は通信利用動向調査、1987〜1995年は携帯電話のサービスが始まった1987年を0として1996年の値まで線でつないだ推計'}
out['goods']['net'] = {'real': {y: net[y] for y in sorted(net) if y <= 2016}, 'est': ramp(1993, 1996, net[1996]), 'src': 'comm', 'note': '1996〜2016年は通信利用動向調査（世帯で誰かがインターネットを使った割合）。1993〜1995年は一般向けの接続サービスが始まった1993年を0とした推計。2017年以降は2016年の値を使う'}
out['goods']['smart'] = {'real': {y: v for y, v in smart.items() if y <= 2013}, 'est': ramp(2008, 2010, smart[2010]), 'src': 'comm', 'note': '2010〜2013年は通信利用動向調査。2008〜2009年は国内で一般向けに広く売られ始めた2008年を0とした推計'}
out['goods']['washlet'] = {'real': {}, 'est': {}, 'estFrom': 1980, 'src': 'consumer', 'note': '1980〜1991年は、家庭向けに広く売られ始めた1980年を0として、消費動向調査の最初の年（1992年）の値まで線でつないだ推計'}

out['problems'] = problems
with open(OUT, 'w') as f: json.dump(out, f, ensure_ascii=False, indent=1, sort_keys=True)
print('保存した：' + os.path.relpath(OUT, ROOT))
print('雪の印 %d・海沿いの印 %d 市区町村、15歳未満人口がない市区町村 %d' % (len(snow), len(coast), len(miss)))
for p in problems: print('・' + p)
